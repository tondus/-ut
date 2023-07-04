import sys 
import os
from root import  ROOT_DIR
import pprint as pp
import pandas as pd
import sys 
import os
from docx import Document
from docx.shared import Inches
from docx.text.paragraph import Paragraph
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import src.lib.fs as fs_legz
import configparser
import codecs
import src.lib.template as template
import shutil
import src.lib.template as t

col_ricef = 'ricef'
col_ut_description = 'ut_description'
col_source_system = 'source_system'
col_target_system = 'target_system'
col_odata = 'odata'
col_proxy = 'proxy'
col_rfc = 'rfc'
col_other1= 'other1'
col_cpi = 'cpi'
col_pi = 'pi'
col_other2 = 'other2'
col_inbound = 'inbound'
col_outbound = 'outbound'
col_synchronous = 'synchronous'
col_asynchronus = 'asynchronus'

def sxx():
    doc = Document('source/OREO_FD_SD_ZPISDI003.docx')
    tables = doc.tables
    index = 0 
    table = tables[8]
    tbl = table._tbl
    cell = table.cell(6, 0)
    xxx = table.row_cells(0)
    cell.text = "xxxxxxxxxxxxxx"
    xx = table.cell(0,3)

def fill_up_dataframe(df,columns):
    result = {}
    exception_col = ['odata','proxy','rfc','other1','cpi','pi','other2','inbound','outbound','synchronous','asynchronus']
    lastest = {}
    for idx , row in df.iterrows():
        for col in columns:
            if col in result:
                values = result[col]
                if not pd.isna(row[col]):
                    lastest[col] = row[col]
                if col not in exception_col:
                    values.append(lastest[col])
                else:
                    values.append(row[col])
                result[col] = values
            else:
                result[col] = [row[col]]
                lastest[col] = row[col]
    
    df_master = pd.DataFrame()
    for key in result.keys():
        df_master[key] = result[key]
    return df_master

def create_structure(df_master,column_headers):
    result = {}
    for idx , row in df_master.iterrows():
        table = {'source_system' : row[col_source_system],
                 'target_system' : row[col_target_system],
                 'ricef' : row[col_ricef],
                 'ut_description' : row[col_ut_description],
                 'odata' :  True if str(row[col_odata]).strip().lower() == "x" else False ,
                 'proxy' :  True if str(row[col_proxy]).strip().lower() == "x" else False ,
                 'rfc' :  True if str(row[col_rfc]).strip().lower() == "x" else False ,
                 'other1':  row[col_other1], 
                 'cpi' :  True if str(row[col_cpi]).strip().lower() == "x" else False ,
                 'pi' :  True if str(row[col_pi]).strip().lower() == "x" else False ,
                 'other2' :  row[col_other2], 
                 'inbound' :  True if str(row[col_inbound]).strip().lower() == "x" else False ,
                 'outbound' :  True if str(row[col_outbound]).strip().lower() == "x" else False ,
                 'synchronous' :  True if str(row[col_synchronous]).strip().lower() == "x" else False ,
                 'asynchronus' :  True if str(row[col_asynchronus]).strip().lower() == "x" else False ,
             }
        key = row[col_ricef]+"_"+row[col_ut_description]
        key = key.replace(" ","_")
        
        if key in result:
            t = result[key]['tables']
            t.append(table)
        else:
            result[key] = {"ricef": row[col_ricef],
                           "tables" : [table]}
    return result
    
def replace_from_onedrive(p):
    for user_file in os.listdir(p):
        local = 'onedrive'
        if not os.path.exists(local):
            os.mkdir(local)
        master_file = p+'\\'+user_file
        dest_file = local + '/' +user_file
      
        if os.path.exists(master_file):
            if os.path.exists(dest_file):
                 os.remove(dest_file) 
            shutil.copyfile(master_file, dest_file)

def get_merged_dataframe():
    base = 'onedrive'
    df_result = pd.DataFrame()
    result = {}
    for user_file in os.listdir(base):
        path = base + '/' + user_file
        if 'filetransfer' not in user_file: 
            df = pd.read_excel(path)
            column_headers = list(df.columns.values)
            df_master= fill_up_dataframe(df,column_headers)
            for idx, row in df_master.iterrows():
                for key in row.keys():
                    if key in result:
                        items = result[key]
                        items.append(row[key])
                    else:
                        result[key] = [row[key]]
    
    for key in result.keys():
        df_result[key] = result[key]
    return df_result


def clear_files_in_path(p):
    for path in os.listdir(p):
        if  os.path.isfile(os.path.join(p, path)):
            full_path = os.path.join(p, path)
            os.remove(full_path)

def clear_onedrive(p):
    for path in os.listdir(p):
        if not os.path.isfile(os.path.join(p, path)):
            shutil.rmtree(os.path.join(p, path))

def main() -> int:
    clear_files_in_path('output/fs')
    config = configparser.ConfigParser()
    config.read('config_fs.ini', encoding='utf-8')
    df = pd.read_excel('source/fs_legacy.xlsx')
    column_headers = list(df.columns.values)
    df_master= fill_up_dataframe(df,column_headers)
    data = create_structure(df_master,column_headers)
    
    for key in data.keys():
        ricefw = data[key]
        print(key)
        fs_legz.create_document('template/OREO_LEGACY_ONLY_TABLE.docx',ricefw,config,key)
                    
if __name__ == '__main__':
    sys.exit(main())