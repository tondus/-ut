import sys
import os
sys.path.append('../../../')
import pprint as pp
import pandas as pd
import sys 
import os
from docx import Document
from docx.shared import Inches
from docx.text.paragraph import Paragraph
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docxtpl import DocxTemplate
import fs_master_template as fs_master
import configparser
import codecs
from root import ROOT_DIR
 
import src.lib.template as t

col_ricef = 'ricef'
col_source_system = 'source_system'
col_target_system = 'target_system'
col_odata = 'odata'
col_proxy = 'proxy'
col_rfc = 'rfc'
col_other_file = 'other_file'
col_cpi = 'cpi'
col_pi = 'pi'
col_other = 'other'
col_inbound = 'inbound'
col_outbound = 'outbound'
col_synchronous = 'synchronous'
col_asynchronus = 'asynchronus'
col_si = 'si'

def sxx():
    doc = Document('source/OREO_FD_SD_ZPISDI003.docx')
    tables = doc.tables
    index = 0 
    table = tables[8]
    tbl = table._tbl
    cell = table.cell(6, 0)
    xxx = table.row_cells(0)
    cell.text = "xxxxxxxxxxxxxx"
    xx = table.cell(0,3)

def fill_up_dataframe(df,columns):
    result = {}
    exception_col = ['odata','proxy','rfc','other_file','cpi','pi','other','inbound','outbound','synchronous','asynchronus']
    lastest = {}
    for idx , row in df.iterrows():
        for col in columns:
            if col in result:
                values = result[col]
                if not pd.isna(row[col]):
                    lastest[col] = row[col]
                if col not in exception_col:
                    values.append(lastest[col])
                else:
                    values.append(row[col])
                result[col] = values
            else:
                result[col] = [row[col]]
                lastest[col] = row[col]
    
    df_master = pd.DataFrame()
    for key in result.keys():
        df_master[key] = result[key]
    return df_master

def create_structure(excel):
    program = {}
    current_program_name = ""
    for (index,col) in  excel.iterrows():
        program_name = col['SAP ECC 6.0\nProgram name']
        program_desc = col['Program description']
        pi_config = col['PI Config']
        sap_master_table_name = col['SAP Master\nTable name']
        if not pd.isna(program_name):
            # read_template_and_write(program_name,program_desc)
            current_program_name = program_name
            program[program_name] = [{'program_name':program_name,
                                      'program_desc':program_desc,
                                      'pi_config':pi_config,
                                      'sap_master_table_name':sap_master_table_name}]            
        else:
            value = program[current_program_name]
            value.append({'program_name':current_program_name,
                          'program_desc':program_desc,
                          'pi_config':pi_config,
                          'sap_master_table_name':sap_master_table_name})
    return program

def create_template_structure(data):
    result = {}
    data_template = {}
    for key in data.keys():
        items = data[key]

        filenames = []
        desc = None
        for i in items:
            filenames.append(i['pi_config'])
            desc = i['program_desc']
        data_template[key] = {'filenames':filenames,
                              'desc':desc}
       
    for ricefw in data_template.keys():
        filenames = data_template[ricefw]
        table = {'source_system' : "SAP S/4HANA",
                 'target_system' : "SAP Master",
                 'odata' :  False ,
                 'proxy' :  False ,
                 'rfc' :  False ,
                 'other1' :  True ,
                 'cpi' :  False ,
                 'pi' :  True ,
                 'other2' :  False ,
                 'inbound' :  False ,
                 'outbound' :  True ,
                 'desc': filenames['desc'],
                 'synchronous' :  False ,
                 'asynchronus' :  True,
                 'filenames' :  filenames['filenames'],
             }
        
        if ricefw in result:
            t = result[ricefw]['tables']
            t.append(table)
        else:
            result[ricefw] = {"ricef": ricefw,
                             "tables" : [table]}
    return result

def main() -> int:
    config = t.get_config()
    oreo = pd.read_excel( ROOT_DIR+'/fs_files/OREO_Master.xlsx', sheet_name='OR-OREO')
    data_master = create_structure(oreo)
    data = create_template_structure(data_master)

    for key in data.keys():
        ricefw = data[key]
        fs_master.create_document( ROOT_DIR+'/fs_files/OREO_MASTER_ONLY_TABLE.docx',ricefw,config,key)
        # break
                    
if __name__ == '__main__':
    sys.exit(main())